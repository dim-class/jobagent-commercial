"""Test configuration.

Environment variables are set *before* any ``app`` import because
``app.core.config.settings`` and ``app.db.session.engine`` are module-level
singletons. That keeps every test on a throwaway SQLite file and a throwaway
copy of the career strategy, so running the suite never touches the developer's
real database or ``config/career_strategy.yaml``.

No test in this suite makes a network call. The OpenAI agent is always patched.
"""

from __future__ import annotations

import os
import shutil
import tempfile
from pathlib import Path

import pytest

_TMP_ROOT = Path(tempfile.mkdtemp(prefix="jobagent-tests-"))
_PROJECT_ROOT = Path(__file__).resolve().parents[2]

_STRATEGY_COPY = _TMP_ROOT / "career_strategy.yaml"
shutil.copyfile(_PROJECT_ROOT / "config" / "career_strategy.yaml", _STRATEGY_COPY)

os.environ["DATABASE_URL"] = f"sqlite:///{(_TMP_ROOT / 'test.db').as_posix()}"
os.environ["CAREER_STRATEGY_PATH"] = str(_STRATEGY_COPY)
os.environ["LOG_LEVEL"] = "WARNING"
os.environ["MAX_ANALYSES_PER_RUN"] = "3"
# Never inherit a real key: tests must not spend money, and the suite must give
# the same result whether or not the developer has a populated project .env.
# An explicit empty value is required - merely popping the variable would let
# pydantic-settings fall through to the real .env file.
os.environ["OPENAI_API_KEY"] = ""
os.environ["OPENAI_MODEL_FAST"] = "test-model-fast"
os.environ["OPENAI_MODEL_SMART"] = "test-model-smart"

from fastapi.testclient import TestClient  # noqa: E402

from app.core.config import get_settings  # noqa: E402
from app.db.base import Base  # noqa: E402
from app.db.session import SessionLocal, engine, init_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models import Resume  # noqa: E402
from app.services.resume_parser import parse_resume  # noqa: E402


@pytest.fixture(scope="session", autouse=True)
def _create_schema():
    init_db()
    yield
    engine.dispose()


@pytest.fixture(autouse=True)
def clean_state():
    """Reset the database and the career strategy between tests.

    Tests that PUT a strategy would otherwise leak their edit into whatever
    runs next, and the strategy is part of the analysis cache key.
    """
    from app.core.career_strategy import load_strategy

    with engine.begin() as conn:
        # ``offers.accepted_revision_id`` and ``offer_revisions.offer_id``
        # reference each other, so the tables have no valid topological order
        # to delete in. Suspending foreign keys for the wipe is a test-harness
        # concern only - the application always runs with them on.
        conn.exec_driver_sql("PRAGMA foreign_keys=OFF")
        try:
            # No ordering needed with foreign keys suspended - and asking for
            # a topological order would warn, since the cycle above has none.
            for table in Base.metadata.tables.values():
                conn.execute(table.delete())
        finally:
            conn.exec_driver_sql("PRAGMA foreign_keys=ON")
    yield
    shutil.copyfile(_PROJECT_ROOT / "config" / "career_strategy.yaml", _STRATEGY_COPY)
    load_strategy(force=True)


@pytest.fixture
def settings():
    return get_settings()


@pytest.fixture
def db():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def client() -> TestClient:
    with TestClient(app) as test_client:
        yield test_client


# --------------------------------------------------------------------------
# document fixtures - generated at runtime so no binaries live in the repo
# --------------------------------------------------------------------------

RESUME_TEXT = """张三
北京 | zhangsan@example.com | 13800138000

个人简介
5年云计算与系统运维经验，熟悉 AWS 与 Linux 环境下的自动化运维。

工作经历
2021.03-至今 某科技公司 云运维工程师
负责 AWS EC2、VPC、ALB 资源的日常维护与成本优化
使用 Terraform 管理基础设施，编写 Python 与 Shell 自动化脚本
2019.06-2021.02 某银行外包 中间件运维工程师
负责 WebSphere、IHS 中间件与 AIX 服务器的部署与调优

项目经历
核心系统上云项目：将 20 套应用从 AIX 迁移至 AWS，负责容量评估与割接验证
监控体系建设：基于 Prometheus 与 Grafana 搭建统一监控告警平台

专业技能
AWS、Linux、Docker、Kubernetes、Terraform、Python、Shell、CI/CD、Git、Prometheus

证书
AWS Certified Solutions Architect - Associate
RHCE

教育经历
2015.09-2019.06 某大学 计算机科学与技术 本科
"""


def _make_pdf_bytes(text: str = RESUME_TEXT) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    # Use a CJK-capable base font so the Chinese text survives the round trip.
    page.insert_textbox(
        pymupdf.Rect(40, 40, 560, 780), text, fontsize=9, fontname="china-ss", lineheight=1.3
    )
    payload = doc.tobytes()
    doc.close()
    return payload


def _make_docx_bytes(text: str = RESUME_TEXT) -> bytes:
    import io

    from docx import Document

    document = Document()
    for line in text.split("\n"):
        document.add_paragraph(line)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "语言能力"
    table.cell(0, 1).text = "中文（母语）/ 英语（CET-6）"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="session")
def pdf_bytes() -> bytes:
    return _make_pdf_bytes()


@pytest.fixture(scope="session")
def docx_bytes() -> bytes:
    return _make_docx_bytes()


@pytest.fixture
def active_resume(db) -> Resume:
    """An active resume parsed from plain text - required by the analyzer."""
    parsed = parse_resume(RESUME_TEXT.encode("utf-8"), "resume.txt")
    resume = Resume(
        filename="resume.txt",
        file_type=parsed.file_type,
        content_hash=parsed.content_hash,
        raw_text=parsed.raw_text,
        parsed_profile_json=parsed.profile,
        is_active=True,
    )
    db.add(resume)
    db.commit()
    db.refresh(resume)
    return resume


SAMPLE_JD = """岗位职责：
1. 负责 AWS 云平台日常运维，管理 EC2、VPC、ALB 等资源；
2. 使用 Terraform 编写基础设施即代码；
3. 维护 Docker 与 Kubernetes 集群；
4. 编写 Python / Shell 自动化脚本。

任职要求：
1. 1-3 年云计算或运维相关经验；
2. 熟悉 Linux 系统与网络基础；
3. 有 CI/CD 经验者优先。
"""

HELPDESK_JD = """岗位职责：
1. 负责公司员工电脑软硬件故障处理、装机与系统重装；
2. 负责打印机、会议室设备维护；
3. 处理 Helpdesk 工单，协助弱电与综合布线。

任职要求：
1. 大专以上学历，1-3 年桌面运维经验；
2. 熟悉 Windows 系统安装与排障。
"""


def make_job_payload(**overrides) -> dict:
    payload = {
        "title": "云计算工程师",
        "company": "示例科技",
        "city": "北京",
        "salary_text": "20k-30k",
        "experience_text": "1-3年",
        "education_text": "本科",
        "raw_description": SAMPLE_JD,
    }
    payload.update(overrides)
    return payload


@pytest.fixture
def job_payload() -> dict:
    return make_job_payload()


# --------------------------------------------------------------------------
# Playwright fixtures - LOCAL HTML ONLY
# --------------------------------------------------------------------------
#
# These drive a real browser so the Locator-based extractor is exercised for
# real, but they only ever load strings from tests/fixtures. No automated test
# may contact zhipin.com or any other live recruitment site.
#
# Headless here is a test-harness choice. The recruitment browser the user
# drives is always headed - see app/services/browser_session.py.

FIXTURES_DIR = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session")
async def _fixture_browser():
    """One headless Chromium for the whole suite.

    Skips (never fails) when Playwright browsers are not installed, so a fresh
    clone can run the rest of the suite before ``playwright install`` has run.
    """
    from playwright.async_api import async_playwright

    async with async_playwright() as pw:
        try:
            browser = await pw.chromium.launch(headless=True)
        except Exception as exc:  # noqa: BLE001
            pytest.skip(
                "Playwright Chromium unavailable - run "
                f"'python -m playwright install chromium' ({type(exc).__name__})"
            )
        try:
            yield browser
        finally:
            await browser.close()


@pytest.fixture
async def browser_page(_fixture_browser):
    """A blank page for loading local fixture HTML. Fresh context per test."""
    context = await _fixture_browser.new_context()
    page = await context.new_page()
    try:
        yield page
    finally:
        await context.close()


@pytest.fixture
async def fixture_page(browser_page):
    """Load fixture HTML (or an inline string) into a page at a chosen URL.

    The URL is faked with a route interception so ``page.url`` looks like a
    real BOSS URL without any network access.
    """

    async def _load(name: str | None = None, *, html: str | None = None, url: str):
        body = html if html is not None else (FIXTURES_DIR / name).read_text(encoding="utf-8")

        async def _handler(route):
            await route.fulfill(status=200, content_type="text/html; charset=utf-8", body=body)

        # Serve only this exact URL; everything else on the host is aborted so
        # no request can escape to the network.
        await browser_page.route("**/*", lambda route: route.abort())
        await browser_page.route(url, _handler)
        await browser_page.goto(url, wait_until="domcontentloaded")
        return browser_page

    return _load
